"""Evidence report assembly + export (decisions Q14, Q22, B5, B12).

Pulls everything captured during a run into one structure, then renders a Word (.docx)
document and a structured (JSON) export. Mirrors the evidence-report spec
(02-EVIDENCE-REPORT-SPEC.md): summary, debrief, overview (+impact +per-type summary),
team, tasks, timeline, plan-gap punch-list, confidence vote.
"""
import io
import json
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor

from . import db
from .scenarios import SCENARIOS, IMPACT_DIMENSIONS, PHASES


def _total_time(started_at, resolved_at):
    try:
        s = datetime.strptime(started_at, "%Y-%m-%d %H:%M:%S")
        r = datetime.strptime(resolved_at, "%Y-%m-%d %H:%M:%S")
        secs = max(int((r - s).total_seconds()), 0)
        h, rem = divmod(secs, 3600)
        m, _ = divmod(rem, 60)
        return f"{h}h {m}m"
    except (TypeError, ValueError):
        return "—"


def assemble(run_id):
    """Gather all run data into a plain dict for rendering/export."""
    run = db.get_run(run_id)
    client = db.get_client(run["client_org_id"])
    scenario = SCENARIOS.get(run["scenario_key"], {})
    overview = db.get_overview(run_id)
    kv = db.get_kv(run_id)
    participants = [dict(p) for p in db.list_participants(run_id)]
    tasks = [dict(t) for t in db.list_tasks(run_id)]
    events = [dict(e) for e in db.list_events(run_id, newest_first=False)]
    votes = [dict(v) for v in db.list_votes(run_id)]
    gaps = [dict(g) for g in db.gaps_for_client(run["client_org_id"])
            if g["status"] == "validated"]

    impact = []
    for key, label in IMPACT_DIMENSIONS:
        rating = kv.get(f"impact_{key}_rating", "")
        expl = kv.get(f"impact_{key}_expl", "")
        if rating or expl:
            impact.append({"dimension": label, "rating": rating, "explanation": expl})

    summary_fields = []
    for fkey, flabel, _kind in scenario.get("summary_fields", []):
        val = kv.get(f"summary_{fkey}", "")
        if val:
            summary_fields.append({"label": flabel, "value": val})

    scores = [v["score"] for v in votes if v["score"] is not None]
    vote_avg = round(sum(scores) / len(scores), 1) if scores else None

    return {
        "client": client["name"],
        "scenario_title": run["scenario_title"],
        "timezone": run["timezone"] if "timezone" in run.keys() else None,
        "started_at": run["started_at"],
        "resolved_at": run["resolved_at"],
        "total_time": _total_time(run["started_at"], run["resolved_at"]),
        "status": run["status"],
        "overview": dict(overview) if overview else {},
        "impact": impact,
        "summary_title": scenario.get("summary_title", "Incident Summary"),
        "summary_fields": summary_fields,
        "debrief": {
            "right": kv.get("debrief_right", ""),
            "wrong": kv.get("debrief_wrong", ""),
            "improve": kv.get("debrief_improve", ""),
        },
        "closing_notes": kv.get("closing_notes", ""),
        "report_status": kv.get("report_status", "draft"),
        "participants": participants,
        "tasks": tasks,
        "events": events,
        "votes": votes,
        "vote_avg": vote_avg,
        "gaps": gaps,
    }


def context_for_ai(d):
    """Compact text the AI closing-notes generator reads."""
    lines = [f"Client: {d['client']}", f"Scenario: {d['scenario_title']}",
             f"Started: {d['started_at']} | Resolved: {d['resolved_at']} | Total: {d['total_time']}"]
    ov = d["overview"]
    if ov:
        for k in ("latest_status", "detection_summary", "how_discovered", "when_discovered",
                  "who_discovered", "impacted", "history"):
            if ov.get(k):
                lines.append(f"{k.replace('_', ' ').title()}: {ov[k]}")
    for it in d["impact"]:
        lines.append(f"Impact — {it['dimension']}: {it['rating']} — {it['explanation']}")
    for sf in d["summary_fields"]:
        lines.append(f"{sf['label']}: {sf['value']}")
    lines.append("\nTimeline:")
    for e in d["events"]:
        who = f" ({e['acting_name']})" if e.get("acting_name") else ""
        lines.append(f"  {e['occurred_at']} [{e['type']}]{who}: {e['description']}")
    if d["gaps"]:
        lines.append("\nPlan gaps identified:")
        for g in d["gaps"]:
            lines.append(f"  - ({g['severity']}) {g['description']}")
    return "\n".join(lines)


def build_structured(d):
    return json.dumps(d, indent=2, default=str)


# ---- Word (.docx) --------------------------------------------------------
def _h(doc, text, level=1):
    doc.add_heading(text, level=level)


def build_docx(d):
    doc = Document()
    title = doc.add_heading("Incident Response Tabletop — Evidence Report", level=0)
    sub = doc.add_paragraph()
    sub.add_run(f"{d['client']} · {d['scenario_title']}").bold = True
    doc.add_paragraph(
        f"Started: {d['started_at']}    Resolved: {d['resolved_at'] or '—'}    "
        f"Total time: {d['total_time']}    Timezone: {d['timezone'] or 'local'}"
    )
    status_p = doc.add_paragraph()
    r = status_p.add_run(f"Report status: {d['report_status'].upper()}")
    r.font.color.rgb = RGBColor(0x15, 0x73, 0x47) if d["report_status"] == "approved" else RGBColor(0x99, 0x66, 0x00)

    # 1. Executive summary / closing notes
    _h(doc, "1. Executive Summary")
    doc.add_paragraph(d["closing_notes"] or "(Closing notes not yet generated.)")

    # 2. Debrief
    _h(doc, "2. Debrief")
    for label, key in (("What went right", "right"), ("What went wrong", "wrong"),
                       ("What can be improved", "improve")):
        _h(doc, label, level=2)
        doc.add_paragraph(d["debrief"].get(key) or "—")

    # 3. Overview
    _h(doc, "3. Overview")
    ov = d["overview"]
    ov_labels = [("latest_status", "Latest status"), ("detection_summary", "Detection summary"),
                 ("how_discovered", "How discovered"), ("when_discovered", "When discovered"),
                 ("who_discovered", "Who discovered"), ("impacted", "Users / assets impacted"),
                 ("history", "Other pertinent history")]
    for k, label in ov_labels:
        if ov.get(k):
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(ov[k])
    if d["impact"]:
        _h(doc, "Impact to Business", level=2)
        t = doc.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Dimension", "Rating", "Explanation"
        for it in d["impact"]:
            c = t.add_row().cells
            c[0].text, c[1].text, c[2].text = it["dimension"], it["rating"], it["explanation"]
    if d["summary_fields"]:
        _h(doc, d["summary_title"], level=2)
        t = doc.add_table(rows=0, cols=2)
        t.style = "Light Grid Accent 1"
        for sf in d["summary_fields"]:
            c = t.add_row().cells
            c[0].text, c[1].text = sf["label"], sf["value"]

    # 4. Team
    _h(doc, "4. Team")
    for p in d["participants"]:
        role = "Observer" if p["role"] == "observer" else (p.get("incident_role") or "Participant")
        doc.add_paragraph(f"{p['full_name']} — {p.get('title') or ''} ({role})", style="List Bullet")

    # 5. Tasks
    _h(doc, "5. Incident Tasks")
    last_phase = None
    for t in d["tasks"]:
        if t["phase"] != last_phase:
            _h(doc, t["phase"], level=2)
            last_phase = t["phase"]
        mark = "☑" if t["status"] == "done" else "☐"
        asg = f" — {t['assignee']}" if t.get("assignee") else ""
        doc.add_paragraph(f"{mark} {t['title']}{asg}", style="List Bullet")

    # 6. Timeline
    _h(doc, "6. Timeline")
    t = doc.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "Time", "Type", "Who", "Description"
    for e in d["events"]:
        c = t.add_row().cells
        c[0].text = e["occurred_at"] or ""
        c[1].text = e["type"]
        c[2].text = e.get("acting_name") or ""
        c[3].text = e["description"] or ""

    # 7. Plan gaps & recommended changes (punch-list)
    _h(doc, "7. Plan Gaps & Recommended Changes")
    if d["gaps"]:
        t = doc.add_table(rows=1, cols=3)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Severity", "Gap", "Recommended change"
        for g in d["gaps"]:
            c = t.add_row().cells
            c[0].text = g["severity"]
            c[1].text = g["description"] or ""
            c[2].text = g["recommended_change"] or ""
    else:
        doc.add_paragraph("No validated plan gaps recorded.")

    # 8. Confidence vote
    _h(doc, "8. Team Confidence Vote (EOS, 1–10)")
    if d["votes"]:
        for v in d["votes"]:
            doc.add_paragraph(f"{v.get('voter') or 'Participant'}: {v['score']}", style="List Bullet")
        if d["vote_avg"] is not None:
            p = doc.add_paragraph()
            p.add_run(f"Average: {d['vote_avg']} / 10").bold = True
    else:
        doc.add_paragraph("No votes recorded.")

    doc.add_paragraph()
    foot = doc.add_paragraph(f"Generated {db.now_str()}.")
    foot.runs[0].font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
